#!/usr/bin/env python3
"""
Convierte archivos Markdown a Word (.docx) manteniendo formato:
- Encabezados (# ## ###)
- Tablas
- Bloques de código
- Listas
- Negrita/itálica
- Reglas horizontales
"""

import re
import markdown
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

class MarkdownToDocx:
    def __init__(self):
        self.doc = Document()
        self.setup_styles()
        
    def setup_styles(self):
        """Configura estilos personalizados para el documento."""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Estilo para código en línea
        if 'InlineCode' not in self.doc.styles:
            inline_code_style = self.doc.styles.add_style('InlineCode', WD_STYLE_TYPE.CHARACTER)
            inline_code_style.font.name = 'Consolas'
            inline_code_style.font.size = Pt(10)
            inline_code_style.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            
        # Estilo para bloques de código
        if 'CodeBlock' not in self.doc.styles:
            code_block_style = self.doc.styles.add_style('CodeBlock', WD_STYLE_TYPE.PARAGRAPH)
            code_block_style.font.name = 'Consolas'
            code_block_style.font.size = Pt(9.5)
            code_block_style.paragraph_format.space_before = Pt(6)
            code_block_style.paragraph_format.space_after = Pt(6)
            code_block_style.paragraph_format.left_indent = Inches(0.3)
            
    def add_heading(self, text, level):
        heading = self.doc.add_heading(text, level=min(level, 3))
        return heading
        
    def add_paragraph_with_formatting(self, text, style=None):
        """Añade párrafo procesando markdown inline (bold, italic, code)."""
        p = self.doc.add_paragraph(style=style)
        self._add_formatted_text(p, text)
        return p
        
    def _add_formatted_text(self, paragraph, text):
        """Procesa texto con formato markdown inline."""
        # Patrones: **bold**, *italic*, `code`
        parts = re.split(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text)
        
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run()
            if part.startswith('**') and part.endswith('**'):
                run.text = part[2:-2]
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run.text = part[1:-1]
                run.italic = True
            elif part.startswith('`') and part.endswith('`'):
                run.text = part[1:-1]
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            else:
                run.text = part
                
    def add_code_block(self, code, language=None):
        """Añade bloque de código con formato monoespaciado."""
        lines = code.split('\n')
        for line in lines:
            p = self.doc.add_paragraph(style='CodeBlock')
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            
    def add_table(self, header_row, data_rows):
        """Añade tabla desde markdown."""
        if not header_row and not data_rows:
            return
            
        # Determinar número máximo de columnas
        max_cols = len(header_row) if header_row else 0
        for row in data_rows:
            max_cols = max(max_cols, len(row))
        if max_cols == 0:
            return
            
        table = self.doc.add_table(rows=1 + len(data_rows), cols=max_cols)
        table.style = 'Light Grid Accent 1'
        
        # Encabezados
        if header_row:
            for i in range(max_cols):
                cell = table.rows[0].cells[i]
                cell.text = ''
                p = cell.paragraphs[0]
                if i < len(header_row):
                    run = p.add_run(header_row[i].strip())
                    run.bold = True
                    run.font.size = Pt(10)
                
        # Datos
        for row_idx, row_data in enumerate(data_rows):
            for col_idx in range(max_cols):
                cell = table.rows[row_idx + 1].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                if col_idx < len(row_data):
                    run = p.add_run(row_data[col_idx].strip())
                    run.font.size = Pt(10)
                
    def add_horizontal_rule(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        # Añadir línea horizontal usando border bottom
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
        
    def add_list_item(self, text, level=0, ordered=False, number=1):
        """Añade item de lista."""
        p = self.doc.add_paragraph(style='List Bullet' if not ordered else 'List Number')
        p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
        self._add_formatted_text(p, text)
        return p

def convert_markdown_to_docx(md_path, docx_path):
    """Convierte un archivo markdown a docx."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    converter = MarkdownToDocx()
    
    # Dividir en líneas para procesar
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_block_lines = []
    code_language = None
    in_table = False
    table_header = None
    table_rows = []
    table_align = None
    
    while i < len(lines):
        line = lines[i]
        
        # Bloque de código
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_language = line[3:].strip() or None
                code_block_lines = []
            else:
                in_code_block = False
                converter.add_code_block('\n'.join(code_block_lines), code_language)
                code_block_lines = []
                code_language = None
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Tabla markdown
        if '|' in line and line.strip().startswith('|'):
            cells = [c.strip() for c in line.strip('|').split('|')]
            
            # Verificar si es fila de alineación (--- | ---)
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                table_align = cells
                i += 1
                continue
                
            if not in_table:
                in_table = True
                table_header = cells
                table_rows = []
            else:
                table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            # Fin de tabla
            converter.add_table(table_header, table_rows)
            in_table = False
            table_header = None
            table_rows = []
            table_align = None
            # No incrementar i, reprocesar esta línea
            continue
            
        # Encabezados
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            converter.add_heading(text, level)
            i += 1
            continue
            
        # Regla horizontal
        if re.match(r'^[-*_]{3,}$', line.strip()):
            converter.add_horizontal_rule()
            i += 1
            continue
            
        # Listas
        list_match = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.+)$', line)
        if list_match:
            indent = len(list_match.group(1))
            marker = list_match.group(2)
            text = list_match.group(3)
            level = indent // 2
            ordered = re.match(r'\d+\.', marker) is not None
            converter.add_list_item(text, level=level, ordered=ordered)
            i += 1
            continue
            
        # Párrafo normal (línea no vacía)
        if line.strip():
            converter.add_paragraph_with_formatting(line)
        else:
            # Línea vacía - añadir espacio
            converter.doc.add_paragraph()
            
        i += 1
        
    # Si terminó en tabla
    if in_table:
        converter.add_table(table_header, table_rows)
        
    # Guardar
    converter.doc.save(docx_path)
    print(f"OK Convertido: {md_path} -> {docx_path}")

if __name__ == '__main__':
    import sys
    if len(sys.argv) != 3:
        print("Uso: python md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert_markdown_to_docx(sys.argv[1], sys.argv[2])